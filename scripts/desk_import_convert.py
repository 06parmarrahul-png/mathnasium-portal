#!/usr/bin/env python3
"""
desk_import_convert.py — the Management Team spreadsheet -> one JSON file.

    python3 scripts/desk_import_convert.py "<file.xlsx>" [out.json]

Produces the payload the Management Desk's import panel expects.

KEEP THE OUTPUT LOCAL. It contains students' and parents' names, so it
must not be committed, emailed, or uploaded anywhere except into Ratio
itself. That is also why the JSON is not in this repository.

WHAT IT DELIBERATELY DOES NOT DO
  Resolve initials to people. That needs the live roster, which only
  exists once somebody is signed in, so it happens in the app at import
  time (src/lib/deskImport.js). The file carries initials as written.

Requires openpyxl:  python3 -m pip install openpyxl
"""
import json
import re
import sys
from datetime import date, datetime

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is needed:  python3 -m pip install openpyxl")

MONTHS = ['january', 'february', 'march', 'april', 'may', 'june',
          'july', 'august', 'september', 'october', 'november', 'december']


def txt(v):
    """A cell as text. A whole number stays a whole number: Excel stores
    invoice numbers as floats, and "581157.0" is not an invoice number."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def as_bool(v):
    return v is True or v == 1 or txt(v).lower() == 'true'


def as_date(v):
    """A real date, or one of the four ways the sheets type them out."""
    if v in (None, ''):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    s = txt(v)
    if re.match(r'^\d{4}-\d{2}-\d{2}', s):
        return s[:10]
    m = re.match(r'^([A-Za-z]+)\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})$', s)
    if m and m.group(1).lower() in MONTHS:
        return '%s-%02d-%02d' % (m.group(3), MONTHS.index(m.group(1).lower()) + 1, int(m.group(2)))
    return None


def as_num(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = txt(v).replace('$', '').replace(',', '')
    try:
        return float(s) if s else None
    except ValueError:
        return None


def status_of(v):
    return 'closed' if txt(v).lower() in ('closed', 'settled', 'complete') else 'open'


def sheet(wb, pattern):
    for ws in wb.worksheets:
        if re.search(pattern, ws.title, re.I):
            return ws
    return None


def rows(ws):
    """Every row below the header, as a dict keyed by column letter."""
    if ws is None:
        return []
    out = []
    for r in range(2, ws.max_row + 1):
        cells = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v not in (None, ''):
                cells[openpyxl.utils.get_column_letter(c)] = v
        if cells:
            out.append(cells)
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit('usage: python3 scripts/desk_import_convert.py "<file.xlsx>" [out.json]')
    src = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else 'desk-import.json'
    wb = openpyxl.load_workbook(src, data_only=True)

    # ── Notes: General and Settled share a column layout ────────────────
    notes = []
    for ws in (sheet(wb, r'^general$'), sheet(wb, r'settled')):
        for r in rows(ws):
            body, subject = txt(r.get('E')), txt(r.get('D'))
            if not body and not subject:
                continue
            notes.append({
                'to': txt(r.get('B')),
                'from': txt(r.get('C')),
                'subject': subject,
                'body': body,
                'loggedAt': as_date(r.get('F')),
                'status': status_of(r.get('H')),
                # Two reply COLUMNS, one per half of the team. They become
                # one thread in the app — the split was a workaround for
                # two people not being able to type in one cell.
                'replies': [t for t in (txt(r.get('G')), txt(r.get('I'))) if t],
            })

    # ── Gift cards ──────────────────────────────────────────────────────
    gift_cards = []
    for r in rows(sheet(wb, r'gift')):
        name = txt(r.get('B'))
        if not name:
            continue
        amount = as_num(r.get('D'))
        gift_cards.append({
            'studentName': name,
            'type': txt(r.get('C')) or 'Other',
            'amount': amount,
            # "Two $5 cards" is not a number. Calling it 5 would be a lie.
            'amountText': '' if amount is not None else txt(r.get('D')),
            'prepurchased': txt(r.get('E')).lower() == 'prepurchased',
            'purchasedOn': as_date(r.get('F')),
            # The sheet's own Status column is NOT carried over: status is
            # derived from this date in the app, and two rows disagreed
            # with their own status.
            'handedOverOn': as_date(r.get('G')),
            'initials': txt(r.get('H')),
            'notes': txt(r.get('I')),
        })

    # ── Referral rally (A–G) and student of the month (I–N), one sheet ──
    rally = rows(sheet(wb, r'referral'))
    referrals = [{
        'studentName': txt(r.get('A')),
        'referredName': txt(r.get('B')),
        'emailSent': as_bool(r.get('C')),
        'prizeCollected': as_bool(r.get('D')),
        'cardsGiven': as_bool(r.get('E')),
        'collectedOn': as_date(r.get('F')),
        'collectedFrom': txt(r.get('G')),
    } for r in rally if txt(r.get('A'))]

    student_of_month = []
    for r in rally:
        month, name = as_date(r.get('I')), txt(r.get('J'))
        if not month or not name or name == '???':
            continue
        student_of_month.append({
            'month': month[:7],
            'studentName': name,
            'writeUp': as_bool(r.get('K')),
            'prizeCollected': as_bool(r.get('L')),
            'onTv': as_bool(r.get('M')),
            'side': 'HS' if txt(r.get('N')).upper().startswith('HS') else 'E',
        })

    # ── Receipts ────────────────────────────────────────────────────────
    receipts = []
    for r in rows(sheet(wb, r'receipt')):
        supplier, description = txt(r.get('A')), txt(r.get('D'))
        if not supplier and not description:
            continue
        receipts.append({
            'supplier': supplier,
            'reference': txt(r.get('B')),
            'orderedOn': as_date(r.get('C')),
            'description': description,
            'amount': as_num(r.get('E')),
            'paymentMethod': txt(r.get('F')),
            'received': bool(re.search(r'receiv', txt(r.get('G')), re.I)),
            # STATUS and Notes were both free text about payment. Joined
            # rather than kept as two columns neither of which could be
            # counted or filtered.
            'notes': ' • '.join(t for t in (txt(r.get('H')), txt(r.get('I'))) if t),
        })

    payload = {
        'source': 'Management Team Post It Notes',
        'generatedAt': datetime.now().isoformat(),
        'notes': notes,
        'giftCards': gift_cards,
        'receipts': receipts,
        'referrals': referrals,
        'studentOfMonth': student_of_month,
    }

    with open(out_path, 'w') as f:
        json.dump(payload, f, indent=1)

    print('wrote', out_path)
    for k in ('notes', 'giftCards', 'receipts', 'referrals', 'studentOfMonth'):
        print('  %-16s %d' % (k, len(payload[k])))
    print('  %-16s %d' % ('notes still open', sum(1 for n in notes if n['status'] == 'open')))


if __name__ == '__main__':
    main()
